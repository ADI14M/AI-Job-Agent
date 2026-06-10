import os
from fpdf import FPDF
from docx import Document
from typing import Dict, Any
from pathlib import Path

BASE_DIR = Path(__file__).resolve().parents[3]
EXPORT_DIR = BASE_DIR / "data" / "exported_reports"
EXPORT_DIR.mkdir(parents=True, exist_ok=True)

class PDF(FPDF):
    def header(self):
        pass

    def footer(self):
        self.set_y(-15)
        self.set_font("Helvetica", "I", 8)
        self.cell(0, 10, f"Page {self.page_no()}", align="C")

def generate_cover_letter_pdf(content: str, filename: str) -> str:
    pdf = PDF()
    pdf.add_page()
    pdf.set_font("Helvetica", size=11)
    
    # Simple multi_cell for text
    pdf.multi_cell(0, 5, content)
    
    filepath = os.path.join(EXPORT_DIR, filename)
    pdf.output(filepath)
    return filepath

def generate_cover_letter_docx(content: str, filename: str) -> str:
    doc = Document()
    doc.add_paragraph(content)
    
    filepath = os.path.join(EXPORT_DIR, filename)
    doc.save(filepath)
    return filepath

def generate_resume_pdf(parsed_data: Dict[str, Any], filename: str) -> str:
    pdf = PDF()
    pdf.add_page()
    
    # Name
    pdf.set_font("Helvetica", "B", 16)
    name = parsed_data.get("name", "Name")
    pdf.cell(0, 10, name, ln=True, align="C")
    
    # Contact Info
    pdf.set_font("Helvetica", size=10)
    contact = f"{parsed_data.get('email', '')} | {parsed_data.get('phone', '')} | {parsed_data.get('location', '')}"
    pdf.cell(0, 5, contact, ln=True, align="C")
    pdf.ln(5)
    
    # Summary
    if parsed_data.get("summary"):
        pdf.set_font("Helvetica", "B", 12)
        pdf.cell(0, 8, "Summary", ln=True)
        pdf.set_font("Helvetica", size=11)
        pdf.multi_cell(0, 5, parsed_data["summary"])
        pdf.ln(5)
        
    # Skills
    if parsed_data.get("skills"):
        pdf.set_font("Helvetica", "B", 12)
        pdf.cell(0, 8, "Skills", ln=True)
        pdf.set_font("Helvetica", size=11)
        skills = ", ".join(parsed_data["skills"])
        pdf.multi_cell(0, 5, skills)
        pdf.ln(5)
        
    # Experience (Simple rendering)
    if parsed_data.get("experience"):
        pdf.set_font("Helvetica", "B", 12)
        pdf.cell(0, 8, "Experience", ln=True)
        for exp in parsed_data["experience"]:
            pdf.set_font("Helvetica", "B", 11)
            pdf.cell(0, 6, f"{exp.get('title', '')} at {exp.get('company', '')}", ln=True)
            pdf.set_font("Helvetica", "I", 10)
            pdf.cell(0, 5, f"{exp.get('start_date', '')} - {exp.get('end_date', '')}", ln=True)
            pdf.set_font("Helvetica", size=11)
            
            for bp in exp.get("bullet_points", []):
                pdf.multi_cell(0, 5, f"* {bp}")
            pdf.ln(3)
            
    filepath = os.path.join(EXPORT_DIR, filename)
    pdf.output(filepath)
    return filepath
